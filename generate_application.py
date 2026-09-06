#!/usr/bin/env python3
"""Generate a court-formatted 恢复强制执行申请书 (.docx)."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


EAST_ASIA_BODY = "仿宋_GB2312"
EAST_ASIA_TITLE = "黑体"
EAST_ASIA_SONG = "宋体"
ASCII_FONT = "Times New Roman"
BODY_SIZE = 12  # 小四
LINE_PT = 22
INDENT_2_CHARS = 24  # 小四两个汉字


def set_run_font(run, *, size_pt, east_asia, bold=False, name=ASCII_FONT):
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), name)


def set_paragraph_format(
    paragraph,
    *,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    first_line_pt=0,
    space_before=0,
    space_after=0,
    line_pt=LINE_PT,
):
    pf = paragraph.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(line_pt)
    pf.widow_control = True
    pf.first_line_indent = Pt(first_line_pt)


def add_text_paragraph(
    doc,
    text,
    *,
    size=BODY_SIZE,
    east_asia=EAST_ASIA_BODY,
    bold=False,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    first_line_pt=INDENT_2_CHARS,
    space_before=0,
    space_after=0,
    line_pt=LINE_PT,
):
    paragraph = doc.add_paragraph()
    set_paragraph_format(
        paragraph,
        align=align,
        first_line_pt=first_line_pt,
        space_before=space_before,
        space_after=space_after,
        line_pt=line_pt,
    )
    run = paragraph.add_run(text)
    set_run_font(run, size_pt=size, east_asia=east_asia, bold=bold)
    return paragraph


def add_mixed_indent_paragraph(doc, label, body):
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph, first_line_pt=INDENT_2_CHARS, line_pt=LINE_PT)
    label_run = paragraph.add_run(label)
    set_run_font(label_run, size_pt=BODY_SIZE, east_asia=EAST_ASIA_SONG, bold=True)
    body_run = paragraph.add_run(body)
    set_run_font(body_run, size_pt=BODY_SIZE, east_asia=EAST_ASIA_BODY, bold=False)
    return paragraph


def set_doc_defaults(doc):
    normal = doc.styles["Normal"]
    normal.font.name = ASCII_FONT
    normal.font.size = Pt(BODY_SIZE)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), ASCII_FONT)
    rfonts.set(qn("w:hAnsi"), ASCII_FONT)
    rfonts.set(qn("w:eastAsia"), EAST_ASIA_BODY)
    rfonts.set(qn("w:cs"), ASCII_FONT)

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)

    styles_el = doc.styles.element
    doc_defaults = styles_el.find(qn("w:docDefaults"))
    if doc_defaults is None:
        return
    rpr_default = doc_defaults.find(qn("w:rPrDefault"))
    if rpr_default is None:
        return
    rpr = rpr_default.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        rpr_default.append(rpr)
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), ASCII_FONT)
    rfonts.set(qn("w:hAnsi"), ASCII_FONT)
    rfonts.set(qn("w:eastAsia"), EAST_ASIA_BODY)


def add_page_number(doc):
    footer = doc.sections[0].footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)

    prefix = paragraph.add_run("— ")
    set_run_font(prefix, size_pt=9, east_asia=EAST_ASIA_SONG)

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    page_run = paragraph.add_run()
    set_run_font(page_run, size_pt=9, east_asia=EAST_ASIA_SONG)
    page_run._r.append(fld_begin)
    page_run._r.append(instr)
    page_run._r.append(fld_sep)
    text_el = OxmlElement("w:t")
    text_el.text = "1"
    page_run._r.append(text_el)
    page_run._r.append(fld_end)

    suffix = paragraph.add_run(" —")
    set_run_font(suffix, size_pt=9, east_asia=EAST_ASIA_SONG)


def build():
    doc = Document()
    set_doc_defaults(doc)
    add_page_number(doc)

    title = doc.add_paragraph()
    set_paragraph_format(
        title,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line_pt=0,
        space_before=0,
        space_after=6,
        line_pt=28,
    )
    title_run = title.add_run("恢复强制执行申请书")
    set_run_font(title_run, size_pt=16, east_asia=EAST_ASIA_TITLE, bold=True)

    parties = [
        (
            "申请人：",
            "崔云飞，男，汉族，2000年10月15日出生，公民身份号码：130528200010151814，"
            "住河北省石家庄市长安区高远森霖城三区13号楼1单元601室，联系电话：17610812519。",
        ),
        (
            "申请人：",
            "于鹤磊，女，汉族，1998年9月29日出生，公民身份号码：230126199809291926，"
            "住河北省石家庄市长安区高远森霖城三区13号楼1单元601室，联系电话：17610812519。",
        ),
        (
            "被申请人：",
            "河北牧銮科技有限公司，住所地：河北省石家庄市裕华区育才街322号一楼，"
            "统一社会信用代码：91130108MA0CLU8T2X。",
        ),
        (
            "法定代表人：",
            "佟耀东，联系电话：13403344424。",
        ),
        (
            "被申请人：",
            "佟耀东，男，1987年8月26日出生，公民身份号码：130826198708260039，"
            "住河北省承德市丰宁满族自治县汤河乡汤河村，联系电话：13403344424。",
        ),
    ]
    for label, body in parties:
        add_mixed_indent_paragraph(doc, label, body)

    add_text_paragraph(
        doc,
        "申请事项",
        east_asia=EAST_ASIA_TITLE,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line_pt=0,
        space_before=6,
        space_after=2,
        line_pt=22,
        size=14,
    )

    items = [
        "一、请求贵院对（2025）冀0108执4502号执行案件恢复强制执行，强制二被申请人立即履行河北省石家庄市裕华区人民法院（2025）冀0108民初4406号民事调解书确定的给付义务，向申请人支付扣除已还款后仍欠付的工资款人民币35000元及相应利息、诉讼费用625元。",
        "二、请求贵院对二被申请人恢复采取限制高消费等强制执行措施。",
        "三、请求贵院通过网络执行查控系统查询二被申请人名下银行存款、微信支付、支付宝及其他财产，并在本案执行标的额范围内予以冻结、扣划。",
        "四、请求贵院调取二被申请人自执行和解协议达成之日起至本申请提出之日止的全部银行及其他支付账户交易流水，核查是否存在转移、隐匿财产行为；一经查实，依法追回相应财产并追究法律责任。",
        "五、本案执行费用，以及二被申请人未按生效法律文书指定期间履行给付金钱义务所应加倍支付的迟延履行期间债务利息，均由二被申请人负担。",
    ]
    for item in items:
        add_text_paragraph(doc, item)

    add_text_paragraph(
        doc,
        "事实与理由",
        east_asia=EAST_ASIA_TITLE,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line_pt=0,
        space_before=6,
        space_after=2,
        line_pt=22,
        size=14,
    )

    facts = [
        "申请人与被申请人河北牧銮科技有限公司、佟耀东追索劳动报酬纠纷一案，经河北省石家庄市裕华区人民法院主持调解，作出（2025）冀0108民初4406号民事调解书，确定二被申请人的给付义务。该调解书已经发生法律效力。因二被申请人未按调解书确定的内容履行义务，申请人依法向贵院申请强制执行，执行案号为（2025）冀0108执4502号。",
        "案件执行过程中，在贵院主持下，申请人与二被申请人达成执行和解协议，当时尚欠工资款人民币43000元及诉讼费用625元。协议约定二被申请人于确定的履行期限内清偿全部剩余债务。基于二被申请人作出的还款承诺，申请人向贵院申请解除了对二被申请人的限制高消费措施。此后，本案以执行和解方式结案并已归档。",
        "和解后，二被申请人并未按约一次性清偿，仅陆续支付部分款项，具体为：2026年1月20日1500元、2026年3月3日1500元、2026年3月21日1500元、2026年4月21日1500元、2026年5月22日1500元、2026年6月27日500元，合计人民币8000元。扣除上述已付款项后，仍欠付工资款人民币35000元及诉讼费用625元。自2026年6月27日最后一笔还款后，二被申请人再未支付任何款项，并以“没钱”为由推诿拖延，未再继续履行，执行和解协议确定的义务至今未能全面履行完毕，有隐匿、转移财产的重大嫌疑。",
        "根据《最高人民法院关于执行和解若干问题的规定》第九条之规定，被执行人一方不履行执行和解协议的，申请执行人可以申请恢复执行原生效法律文书。依据《中华人民共和国民事诉讼法》第二百六十四条之规定，被执行人未按法律文书指定的期间履行给付金钱义务的，应当加倍支付迟延履行期间的债务利息。同时，依照《最高人民法院关于民事执行中财产调查若干问题的规定》《最高人民法院关于限制被执行人高消费及有关消费的若干规定》等规定，申请执行人有权请求人民法院通过网络执行查控系统查询、冻结被执行人财产，并对被执行人采取限制高消费等强制措施。",
        "为维护申请人的合法权益，特依法向贵院提出恢复强制执行申请，请予审查准许。",
    ]
    for para in facts:
        add_text_paragraph(doc, para)

    add_text_paragraph(
        doc,
        "此致",
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_line_pt=INDENT_2_CHARS,
        space_before=6,
    )
    add_text_paragraph(
        doc,
        "河北省石家庄市裕华区人民法院",
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_line_pt=0,
    )

    add_text_paragraph(
        doc,
        "申请人（签字）：崔云飞",
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        first_line_pt=0,
        space_before=12,
    )
    add_text_paragraph(
        doc,
        "申请人（签字）：于鹤磊",
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        first_line_pt=0,
        space_before=4,
    )
    add_text_paragraph(
        doc,
        "2026年9月6日",
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        first_line_pt=0,
        space_before=4,
    )

    add_text_paragraph(
        doc,
        "附：",
        east_asia=EAST_ASIA_SONG,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_line_pt=0,
        space_before=10,
        line_pt=20,
        size=10.5,
    )
    attachments = [
        "1. 申请人身份证复印件；",
        "2. （2025）冀0108民初4406号民事调解书复印件；",
        "3. （2025）冀0108执4502号执行案件相关材料；",
        "4. 执行和解协议复印件；",
        "5. 已还款记录（2026年1月20日至2026年6月27日，合计8000元）；",
        "6. 其他证明材料。",
    ]
    for att in attachments:
        add_text_paragraph(
            doc,
            att,
            align=WD_ALIGN_PARAGRAPH.LEFT,
            first_line_pt=0,
            line_pt=20,
            size=10.5,
        )

    out = "/workspace/恢复强制执行申请书.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    build()
